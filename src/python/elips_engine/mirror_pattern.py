# elips — Editor de LIcitacions PúbliqueS
# Copyright (C) 2026  Francesc Rambla i Marigot
#
# This program is free software: you can redistribute it and/or modify
# it under the terms of the GNU General Public License as published by
# the Free Software Foundation, either version 3 of the License, or
# (at your option) any later version.
#
# This program is distributed in the hope that it will be useful,
# but WITHOUT ANY WARRANTY; without even the implied warranty of
# MERCHANTABILITY or FITNESS FOR A PARTICULAR PURPOSE.  See the
# GNU General Public License for more details.
#
# You should have received a copy of the GNU General Public License
# along with this program.  If not, see <https://www.gnu.org/licenses/>.

"""Detects and replicates a "mirror" column pattern: an OUT_ sheet whose
existing columns are all simple `=SourceSheet!Cell` link formulas pointing
at another sheet, so a newly-added field can be offered the same treatment
instead of being left unlinked to the sheet that actually holds the data."""

import json
import re

from openpyxl import load_workbook

from .text_utils import sanitize_id, _col_letter, _col_index_from_letter
from .excel_io import is_simple_link_formula, REF_REGEX


def _find_sheet_for_group(wb, sheet_name):
    """Resolves a schema group name (e.g. 'nomconjunt' or 'pres.parts') to the
    actual sheet in `wb`, tolerating the OUT_/JSON_/EXPORT_ prefixes the rest
    of the engine uses. The OUT_-prefixed sheet (if any) wins over a bare
    same-named sheet, since callers always pass the app's schema group name —
    which is exactly the OUT_ sheet's name with the prefix stripped — and a
    same-named *source* sheet existing alongside it (the mirror scenario this
    function primarily exists for) must never shadow it."""
    for pfx in ('OUT_', 'JSON_', 'EXPORT_'):
        candidate = f"{pfx}{sheet_name}"
        if candidate in wb.sheetnames:
            return candidate
    if sheet_name in wb.sheetnames:
        return sheet_name
    clean_target = sanitize_id(sheet_name, allow_dots=True)
    for s in wb.sheetnames:
        stripped = s
        for pfx in ('OUT_', 'JSON_', 'EXPORT_'):
            if s.upper().startswith(pfx):
                stripped = s[len(pfx):]
                break
        if sanitize_id(stripped, allow_dots=True) == clean_target:
            return s
    return None


def analyze_mirror_pattern(excel_path, sheet_name):
    """
    Scans a tabular sheet's EXISTING columns to determine whether the sheet
    is a simple, cell-by-cell mirror of another sheet — every non-empty cell
    in every existing column is a plain `=SourceSheet!Cell` link formula (see
    `is_simple_link_formula`), all columns point at the SAME single source
    sheet, each OUT_ column maps to exactly one source column, and every row
    maps to the source with the same constant row offset ("one column per
    source column", so a new column could plausibly be appended the same way).

    Used by the "add a new field/column to a group" flow (GroupConfigModal.vue
    via useGroupMetadata.js's saveGroupConfig) to decide whether to offer
    replicating a newly-added column into the source sheet too, instead of
    silently leaving the new OUT_ column unlinked to whatever sheet actually
    holds the real data (see bug report: OUT_nomconjunt mirroring `nomconjunt`).

    Returns a JSON string:
      {"is_mirror": true, "source_sheet": "...", "row_offset": N,
       "next_source_col": "E", "header_row": 1, "first_data_row": 2,
       "last_data_row": N}
    or
      {"is_mirror": false, "reason": "..."}
    reason is one of: sheet_not_found, no_data_rows, no_formulas_found,
    non_formula_or_complex_cell, unparseable_formula, unparseable_cell_ref,
    multiple_or_no_source_sheets, column_maps_to_multiple_source_columns,
    inconsistent_row_offset, duplicate_source_column_mapping.
    """
    try:
        wb = load_workbook(excel_path)
        target_name = _find_sheet_for_group(wb, sheet_name)
        if not target_name:
            return json.dumps({"is_mirror": False, "reason": "sheet_not_found"})

        ws = wb[target_name]
        if ws.max_row < 2 or ws.max_column < 1:
            return json.dumps({"is_mirror": False, "reason": "no_data_rows"})

        # First pass: only check whether the sheet has ANY formula at all. A
        # sheet with none is simply not attempting to mirror anything — the
        # ordinary case for every brand-new/plain-data sheet — and must be
        # told apart from one that mixes formulas with plain values (a
        # genuinely ambiguous, only-partially-mirrored sheet).
        any_formula = False
        for col_idx in range(1, ws.max_column + 1):
            if ws.cell(1, col_idx).value in (None, ''):
                continue
            for row_idx in range(2, ws.max_row + 1):
                val_str = str(ws.cell(row_idx, col_idx).value or '').strip()
                if val_str and is_simple_link_formula(val_str):
                    any_formula = True
                    break
            if any_formula:
                break
        if not any_formula:
            return json.dumps({"is_mirror": False, "reason": "no_formulas_found"})

        source_sheets = set()
        col_source_cols = {}
        row_offsets = set()

        for col_idx in range(1, ws.max_column + 1):
            header_val = ws.cell(1, col_idx).value
            if header_val in (None, ''):
                continue
            for row_idx in range(2, ws.max_row + 1):
                cell = ws.cell(row_idx, col_idx)
                val_str = str(cell.value if cell.value is not None else '').strip()
                if val_str == '':
                    continue
                if not is_simple_link_formula(val_str):
                    return json.dumps({
                        "is_mirror": False,
                        "reason": "non_formula_or_complex_cell",
                        "cell": f"{_col_letter(col_idx)}{row_idx}"
                    })
                m = REF_REGEX.search(val_str)
                if not m:
                    return json.dumps({"is_mirror": False, "reason": "unparseable_formula"})
                source_sheets.add(m.group(1) or m.group(2))
                cell_coord = m.group(3).replace('$', '')
                cm = re.match(r'^([A-Za-z]+)([0-9]+)$', cell_coord)
                if not cm:
                    return json.dumps({"is_mirror": False, "reason": "unparseable_cell_ref"})
                col_source_cols.setdefault(col_idx, set()).add(cm.group(1).upper())
                row_offsets.add(int(cm.group(2)) - row_idx)
        if len(source_sheets) != 1:
            return json.dumps({
                "is_mirror": False,
                "reason": "multiple_or_no_source_sheets",
                "sheets": sorted(source_sheets)
            })
        if any(len(cols) != 1 for cols in col_source_cols.values()):
            return json.dumps({"is_mirror": False, "reason": "column_maps_to_multiple_source_columns"})
        if len(row_offsets) != 1:
            return json.dumps({"is_mirror": False, "reason": "inconsistent_row_offset"})

        used_col_nums = [_col_index_from_letter(next(iter(cols))) for cols in col_source_cols.values()]
        if len(used_col_nums) != len(set(used_col_nums)):
            return json.dumps({"is_mirror": False, "reason": "duplicate_source_column_mapping"})

        return json.dumps({
            "is_mirror": True,
            "source_sheet": next(iter(source_sheets)),
            "row_offset": next(iter(row_offsets)),
            "next_source_col": _col_letter(max(used_col_nums) + 1),
            "header_row": 1,
            "first_data_row": 2,
            "last_data_row": ws.max_row,
        })
    except Exception as ex:
        return json.dumps({"is_mirror": False, "reason": f"error: {ex}"})


def apply_mirror_column(excel_path, sheet_name, new_field_name, out_excel_path):
    """
    Given a sheet already confirmed (by the caller, after `analyze_mirror_pattern`
    returned is_mirror=true AND the user explicitly approved it) to be a simple
    mirror of another sheet, appends `new_field_name` as a new column in BOTH
    the source sheet (a plain header cell — its actual data is left for the
    user to fill in directly, same as any other new column in that sheet) and
    the OUT_ sheet (a `=SourceSheet!Col<row>` formula per existing data row,
    replicating the exact convention `analyze_mirror_pattern` detected).
    Re-runs the analysis itself rather than trusting pre-computed coordinates
    from the caller, so this is safe to call on its own.
    Returns a JSON string: {"applied": true, ...} or {"applied": false, "reason": ...}.
    """
    try:
        analysis = json.loads(analyze_mirror_pattern(excel_path, sheet_name))
        if not analysis.get("is_mirror"):
            return json.dumps({"applied": False, "reason": analysis.get("reason", "not_a_mirror")})

        wb = load_workbook(excel_path)
        target_name = _find_sheet_for_group(wb, sheet_name)
        # analysis["source_sheet"] came straight out of the formula text
        # (REF_REGEX), so it is already the exact sheet name — no OUT_-prefix
        # guessing needed (and none wanted: that heuristic exists to resolve
        # schema group names, and would wrongly prefer an OUT_ sheet here).
        source_name = analysis["source_sheet"]
        if source_name not in wb.sheetnames:
            return json.dumps({"applied": False, "reason": "source_sheet_missing"})

        ws = wb[target_name]
        ws_source = wb[source_name]

        new_col_idx = ws.max_column + 1
        for c in range(1, ws.max_column + 1):
            if str(ws.cell(1, c).value or '').strip() == str(new_field_name).strip():
                return json.dumps({"applied": False, "reason": "column_already_exists"})

        source_col_letter = analysis["next_source_col"]
        source_col_idx = _col_index_from_letter(source_col_letter)
        existing_source_header = ws_source.cell(1, source_col_idx).value
        if existing_source_header not in (None, ''):
            return json.dumps({"applied": False, "reason": "source_column_already_occupied"})

        ws.cell(1, new_col_idx).value = new_field_name
        ws_source.cell(1, source_col_idx).value = new_field_name

        row_offset = analysis["row_offset"]
        for row_idx in range(analysis["first_data_row"], analysis["last_data_row"] + 1):
            ws.cell(row_idx, new_col_idx).value = f"='{source_name}'!{source_col_letter}{row_idx + row_offset}"

        wb.save(out_excel_path)
        return json.dumps({
            "applied": True,
            "source_sheet": source_name,
            "source_col": source_col_letter,
            "out_col": _col_letter(new_col_idx),
        })
    except Exception as ex:
        return json.dumps({"applied": False, "reason": f"error: {ex}"})
